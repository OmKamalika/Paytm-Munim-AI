"""Build PayTMunim_Synthetic_Data.xlsx from generate.py output. Derived fields are native Excel formulas."""
import pickle, sys
from datetime import date, datetime
import openpyxl.workbook.workbook as owb
from openpyxl.styles import Font, PatternFill, Alignment
owb.DEFAULT_FONT = Font(name="Arial", sz=10)
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell

OUT = sys.argv[1] if len(sys.argv) > 1 else "out"
D = pickle.load(open(f"{OUT}/_data.pkl", "rb"))
wb = Workbook(write_only=True)

NAVY, TEAL = "002E6E", "00739E"
F_HDR = Font(name="Arial", sz=10, bold=True, color="FFFFFF")
FILL_RAW, FILL_F = PatternFill("solid", fgColor=NAVY), PatternFill("solid", fgColor=TEAL)
F_IN = Font(name="Arial", sz=10, color="0000FF")
F_B = Font(name="Arial", sz=10, bold=True)
F_T = Font(name="Arial", sz=14, bold=True, color=NAVY)
F_NOTE = Font(name="Arial", sz=9, italic=True, color="555555")
INR, INR2, PCT, DT, DTM = "#,##0;-#,##0;-", "#,##0.00;-#,##0.00;-", "0.0%", "yyyy-mm-dd", "yyyy-mm-dd hh:mm:ss"

def C(ws, v, font=None, fmt=None, fill=None, align=None):
    c = WriteOnlyCell(ws, value=v)
    if font: c.font = font
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    if align: c.alignment = align
    return c

def sheet(name, cols, widths=None, formula_cols=()):
    ws = wb.create_sheet(name)
    for i, w in enumerate(widths or [16] * len(cols)):
        ws.column_dimensions[chr(65 + i) if i < 26 else "A" + chr(65 + i - 26)].width = w
    ws.freeze_panes = "A2"
    ws.append([C(ws, h, F_HDR, fill=FILL_F if i in formula_cols else FILL_RAW) for i, h in enumerate(cols)])
    return ws

# ------------------------------------------------------------------ README
ws = wb.create_sheet("README"); ws.column_dimensions["A"].width = 24; ws.column_dimensions["B"].width = 110
for r in [
    [C(ws, "PayTMunim · synthetic data set (seed 42, generator v1)", F_T)],
    [C(ws, "ALL DATA IS SYNTHETIC. Names, accounts and amounts are illustrative. No real person or business.", F_B)],
    [],
    [C(ws, "Colour key", F_B), "Navy header = raw data (what the source system would send). Teal header = derived with an Excel formula (the logic the backend must reproduce). Blue text = an input you can change."],
    [C(ws, "Business", F_B), "Sharma General Stores · owner Rakesh Sharma · 3 outlets (Dadar, Matunga, Sion) · one Paytm MID · one bank account ··4029"],
    [C(ws, "Period", F_B), "Paytm and handovers 1 May 2025 to 30 Oct 2026 (30 Oct is 'today', Paytm until 12:40 pm, no handovers yet). Bank statement to close of 29 Oct 2026."],
    [],
    [C(ws, "Sheet", F_B), C(ws, "What it holds / key logic", F_B)],
    ["Assumptions", "Every changeable input: MDR, GST on MDR, credit-cap %, calibration window, forecast inflows, drawdown amount and date."],
    ["Outlets, DeviceMap, OtherQRs, Staff", "Masters. DeviceMap is effective-dated: QR PQR-1002 moved Matunga → Sion on 15 Oct 2026 (PRD 14 AC-03)."],
    ["Paytm_Txns", "One row per Paytm payment or refund (~1.15 lakh). Paytm sends device_id, not outlet: outlet_id is derived by an effective-dated lookup on DeviceMap. Also txn_date, shift, MDR (cards 1.8%), GST on MDR, settlement date (T+1)."],
    ["Paytm_Settlements", "Expected Paytm payout per day (net collections − MDR − GST) vs the bank credit on T+1. MATCH / CHECK / PENDING."],
    ["Bank_Statement", "Account Aggregator (ReBIT DEPOSIT) field names. transactionalBalance is a running formula from the opening balance. label_* columns are ground truth for testing outflow detection (not in real AA data)."],
    ["Shift_Handovers", "Per outlet, per shift: system Paytm total, other-QR totals, cash counted, cash spent from the till. collected_total is a formula. Matunga Thursday evenings: 3 of the last 4 are short (PRD 19)."],
    ["Suppliers, Invoices", "6 approved suppliers + 1 unknown. Invoices include GD-1182 and three attack variants (duplicate, injected text, changed account)."],
    ["Upcoming_Outflows", "The 7 known outflows 31 Oct to 20 Nov 2026 (total ₹12,66,500)."],
    ["Lenders, Existing_EMIs, Owner_Mandate", "5 lenders with match logic, EMI via PMT, EMIs-after vs credit cap. Mandate thresholds for Intent Guard rules."],
    ["GST_Returns", "18 monthly GSTR-1 / GSTR-3B summaries. Output tax and net payable are formulas."],
    ["Monthly_Summary", "Per outlet per month: Paytm net, all-channel collections, Paytm share."],
    ["Recurring_Bills", "Recurring debits found from statement narrations (EMI, electricity, broadband, rent, salaries, GST), with amount, typical day and next due date."],
    ["Bill_Uploads", "How open bills arrive: owner or staff upload via WhatsApp, app or voice. Unpaid balances come from these, not from bank data."],
    ["Cash_Plan", "30 Oct to 20 Nov: opening ₹4.5 lakh, forecast inflows, known outflows, recurring bills → lowest point −₹1,84,500 on 5 Nov; with a ₹2 lakh drawdown on 4 Nov, +₹15,500."],
    ["Calibration", "PRD 01 targets checked by formula against the data: PASS / FAIL."],
    [],
    [C(ws, "CSV copies", F_B), "csv/*.csv hold the same tables as values (derived fields pre-computed) for the Postgres loader. Re-run generate.py with seed 42 to reproduce them byte for byte."],
]:
    ws.append(r)

# ------------------------------------------------------------------ Assumptions
ws = wb.create_sheet("Assumptions")
for col, w in zip("ABC", (40, 18, 70)): ws.column_dimensions[col].width = w
A = [  # row 2.. : label, value, fmt, note
    ("Card MDR (cards on Paytm EDC)", 0.018, PCT, "Team assumption; UPI MDR = 0"),
    ("GST on MDR", 0.18, PCT, "18% GST on payment charges"),
    ("Credit cap, % of average monthly collections", 0.08, PCT, "Owner mandate (PRD 07 R11)"),
    ("Bank opening balance, 1 May 2025", D["OPENING"], INR2, "Generator output"),
    ("Cash in hand across tills, close of 29 Oct 2026", D["CASH"], INR, "Generator output: undeposited cash + ₹10,000 float per till"),
    ("Drawdown amount (option 2)", 200000, INR, "Cash-plan option 2"),
    ("Drawdown date", date(2026, 11, 4), DT, "Cash-plan option 2"),
    ("Calibration window from", date(2025, 10, 1), DT, "12 complete months"),
    ("Calibration window to", date(2026, 9, 30), DT, ""),
    ("Today", date(2026, 10, 30), DT, "Demo 'today'"),
    ("Forecast combined inflow a day, 30 Oct to 3 Nov", 70000, INR, "PRD 03 forecast output (fixed for the demo)"),
    ("Forecast combined inflow a day, from 4 Nov", 120000, INR, "PRD 03 forecast output, Diwali week"),
    ("Forecast step-up date", date(2026, 11, 4), DT, ""),
    ("Monthly sales tolerance", 0.02, PCT, "PRD 01 calibration"),
    ("Diwali uplift tolerance (×)", 0.05, "0.00", ""),
    ("Paytm share target", 0.25, PCT, ""),
    ("Paytm share tolerance (points)", 0.03, PCT, ""),
    ("Diwali week 2025 from", date(2025, 10, 15), DT, ""),
    ("Diwali week 2025 to", date(2025, 10, 21), DT, ""),
    ("Reference weeks from (4 normal weeks)", date(2025, 9, 8), DT, ""),
    ("Reference weeks to", date(2025, 10, 5), DT, ""),
    ("Shift range band (±)", 0.04, PCT, "PRD 19: expected ± 12-week spread"),
    ("Drawdown interest rate (Lender A)", 0.185, PCT, "Lender A"),
    ("Early repayment date", date(2026, 11, 30), DT, "Repay drawdown by 30 Nov"),
]
ws.append([C(ws, "Input", F_HDR, fill=FILL_RAW), C(ws, "Value", F_HDR, fill=FILL_RAW), C(ws, "Source / note", F_HDR, fill=FILL_RAW)])
for lab, v, fmt, note in A:
    ws.append([lab, C(ws, v, F_IN, fmt), C(ws, note, F_NOTE)])
AS = {lab: f"Assumptions!$B${i + 2}" for i, (lab, *_r) in enumerate(A)}
a = lambda k: AS[k]

# ------------------------------------------------------------------ masters
ws = sheet("Business", list(D["BUSINESS"]), [22] * len(D["BUSINESS"])); ws.append(list(D["BUSINESS"].values()))
ws = sheet("Outlets", ["outlet_id", "name", "open_time", "close_time", "shifts", "monthly_target_inr"], [10, 12, 10, 10, 30, 18])
for o, n, t in D["OUTLETS"]: ws.append([o, n, "08:00", "22:00", "Morning 08-15; Evening 15-22", C(ws, t, None, INR)])
ws = sheet("DeviceMap", ["device_id", "device_type", "serial_end", "outlet_id", "valid_from", "valid_to"], [12, 14, 10, 10, 12, 12])
for r in D["DEVICE_MAP"]: ws.append([r[0], r[1], r[2], r[3], C(ws, r[4], None, DT), C(ws, r[5], None, DT)])
NDM = len(D["DEVICE_MAP"]) + 1
ws = sheet("OtherQRs", ["qr_id", "label", "share_of_sales"], [8, 24, 14])
for q in D["OTHER_QRS"]: ws.append([q[0], q[1], C(ws, q[2], None, PCT)])
ws = sheet("Staff", ["staff_id", "name", "role", "outlet_id", "shift", "monthly_salary_inr"], [10, 16, 10, 10, 10, 18])
for s in D["STAFF"]: ws.append(list(s[:5]) + [C(ws, s[5], None, INR)])

# ------------------------------------------------------------------ Paytm_Txns
P = D["paytm"]; NP = len(P) + 1
ws = sheet("Paytm_Txns", ["txn_id", "txn_timestamp", "txn_date", "shift", "mid", "device_id", "outlet_id", "method",
                          "txn_type", "amount_inr", "orig_txn_id", "payer_vpa_hash", "mdr_inr", "gst_on_mdr_inr", "settlement_date"],
           [13, 19, 11, 9, 13, 10, 11, 9, 9, 11, 13, 17, 9, 11, 12], formula_cols=(2, 3, 6, 12, 13, 14))
mid = D["BUSINESS"]["paytm_mid"]
for i, r in enumerate(P, start=2):
    ws.append([r["txn_id"], C(ws, r["ts"], None, DTM), C(ws, f"=INT(B{i})", None, DT),
               f'=IF(HOUR(B{i})<15,"Morning","Evening")', mid, r["device"],
               f'=IFERROR(INDEX(DeviceMap!$D$2:$D${NDM},MATCH(1,INDEX((DeviceMap!$A$2:$A${NDM}=F{i})*(DeviceMap!$E$2:$E${NDM}<=C{i})*(DeviceMap!$F$2:$F${NDM}>=C{i}),0),0)),"UNASSIGNED")',
               r["method"], r["type"], r["amount"], r["orig"] or None, r["payer"],
               C(ws, f'=IF(AND(H{i}="CARD",I{i}="PAYMENT"),ROUND(ABS(J{i})*{a("Card MDR (cards on Paytm EDC)")},2),0)', None, INR2),
               C(ws, f"=ROUND(M{i}*{a('GST on MDR')},2)", None, INR2), C(ws, f"=C{i}+1", None, DT)])
PT = lambda col: f"Paytm_Txns!${col}$2:${col}${NP}"

# ------------------------------------------------------------------ Bank_Statement
B = D["bank"]; NB = len(B) + 1
ws = sheet("Bank_Statement", ["txnId", "type", "mode", "amount", "transactionalBalance", "transactionTimestamp", "valueDate", "narration",
                              "reference", "label_category", "label_payee", "label_outlet"],
           [11, 8, 8, 13, 14, 19, 11, 62, 16, 18, 22, 10], formula_cols=(4, 6))
for i, b in enumerate(B, start=2):
    prev = a("Bank opening balance, 1 May 2025") if i == 2 else f"E{i - 1}"
    ws.append([b["txn_id"], b["type"], b["mode"], C(ws, b["amount"], None, INR2),
               C(ws, f'=ROUND({prev}+IF(B{i}="CREDIT",D{i},-D{i}),2)', None, INR2), C(ws, b["ts"], None, DTM),
               C(ws, f"=INT(F{i})", None, DT), b["narr"] or None, b["ref"], b["cat"], b["payee"] or None, b["outlet"] or None])
BS = lambda col: f"Bank_Statement!${col}$2:${col}${NB}"
LAST_BAL = f"Bank_Statement!$E${NB}"

# ------------------------------------------------------------------ Shift_Handovers
H = D["hand"]; NH = len(H) + 1
ws = sheet("Shift_Handovers", ["handover_id", "date", "outlet_id", "shift", "staff_id", "status", "paytm_total_inr", "qr2_total_inr",
                               "qr3_total_inr", "cash_count_inr", "cash_expenses_till_inr", "expense_note", "collected_total_inr"],
           [11, 11, 9, 9, 9, 13, 14, 12, 12, 13, 15, 12, 16], formula_cols=(12,))
for i, h in enumerate(H, start=2):
    ws.append([f"HND{i - 1:05d}", C(ws, h["date"], None, DT), h["outlet"], h["shift"], h["staff"], h["status"],
               C(ws, h["paytm"], None, INR), C(ws, h["qr2"], None, INR), C(ws, h["qr3"], None, INR), C(ws, h["cash_count"], None, INR),
               C(ws, h["exp_till"], None, INR), h["exp_note"] or None, C(ws, f"=G{i}+H{i}+I{i}+J{i}+K{i}", None, INR)])
SH = lambda col: f"Shift_Handovers!${col}$2:${col}${NH}"

# ------------------------------------------------------------------ Paytm_Settlements
ws = sheet("Paytm_Settlements", ["sales_date", "settlement_date", "net_collections_inr", "mdr_inr", "gst_on_mdr_inr",
                                 "expected_payout_inr", "bank_credit_inr", "difference_inr", "status"],
           [11, 14, 18, 11, 13, 18, 15, 13, 10], formula_cols=(1, 2, 3, 4, 5, 6, 7, 8))
d0 = date(2025, 5, 1); days = (date(2026, 10, 30) - d0).days + 1
from datetime import timedelta
for k in range(days):
    i = k + 2; d = d0 + timedelta(k)
    ws.append([C(ws, d, None, DT), C(ws, f"=A{i}+1", None, DT),
               C(ws, f"=SUMIFS({PT('J')},{PT('C')},A{i})", None, INR), C(ws, f"=SUMIFS({PT('M')},{PT('C')},A{i})", None, INR2),
               C(ws, f"=SUMIFS({PT('N')},{PT('C')},A{i})", None, INR2), C(ws, f"=ROUND(C{i}-D{i}-E{i},2)", None, INR2),
               C(ws, f'=SUMIFS({BS("D")},{BS("J")},"paytm_settlement",{BS("G")},B{i})', None, INR2),
               C(ws, f"=ROUND(G{i}-F{i},2)", None, INR2),
               f'=IF(B{i}>MAX({BS("G")}),"PENDING",IF(ABS(H{i})<0.5,"MATCH","CHECK"))'])
NSET = days + 1

# ------------------------------------------------------------------ Suppliers, Invoices, Outflows
ws = sheet("Suppliers", ["supplier_id", "name", "category", "verified_account_last4", "account_changed_on", "approved"], [11, 24, 20, 20, 18, 10])
for s in D["SUPPLIERS"]:
    ws.append([s[0], s[1], s[2], s[4], C(ws, date.fromisoformat(s[5]), None, DT), s[6]])
ws = sheet("Invoices", ["invoice_row", "invoice_no", "supplier_name", "issue_date", "due_date", "amount_inr", "pay_to_account_last4",
                        "received_via", "scenario", "description", "expected_guard_result", "account_matches_verified"],
           [10, 10, 24, 11, 11, 12, 18, 14, 15, 60, 34, 20], formula_cols=(11,))
for i, v in enumerate(D["INVOICES"], start=2):
    ws.append([v[0], v[1], v[2], C(ws, v[3], None, DT), C(ws, v[4], None, DT), C(ws, v[5], None, INR), v[6], v[7], v[8], v[9],
               v[10] or None,
               f'=IFERROR(IF(INDEX(Suppliers!$D$2:$D$8,MATCH(C{i},Suppliers!$B$2:$B$8,0))=G{i},"YES","NO"),"UNKNOWN PAYEE")'])
ws = sheet("Upcoming_Outflows", ["outflow_id", "due_date", "payee", "amount_inr", "category", "invoice_no", "note"], [10, 11, 22, 12, 16, 10, 18])
for r in D["OUTFLOWS"]:
    ws.append([r[0], C(ws, r[1], None, DT), r[2], C(ws, r[3], None, INR), r[4], r[5] or None, r[6]])
ws.append([]); ws.append(["Total", None, None, C(ws, "=SUM(D2:D8)", F_B, INR)])

# ------------------------------------------------------------------ GST
ws = sheet("GST_Returns", ["period", "taxable_0pct", "taxable_5pct", "taxable_12pct", "taxable_18pct", "output_tax", "itc", "net_payable",
                           "gstr1_filed_on", "gstr3b_due_on"], [9, 13, 13, 13, 13, 12, 12, 12, 14, 14], formula_cols=(5, 7))
for i, g in enumerate(D["gst"], start=2):
    ws.append([g["period"], C(ws, g["tv0"], None, INR), C(ws, g["tv5"], None, INR), C(ws, g["tv12"], None, INR), C(ws, g["tv18"], None, INR),
               C(ws, f"=ROUND(C{i}*0.05,0)+ROUND(D{i}*0.12,0)+ROUND(E{i}*0.18,0)", None, INR), C(ws, g["itc"], None, INR),
               C(ws, f"=F{i}-G{i}", None, INR), C(ws, g["gstr1_filed"], None, DT), C(ws, g["gstr3b_due"], None, DT)])
NG = len(D["gst"]) + 1

# ------------------------------------------------------------------ Lenders, EMIs, Mandate
ws = sheet("Existing_EMIs", ["loan_id", "lender", "emi_inr", "debit_day", "ends_on"], [14, 20, 10, 10, 11])
ws.append(["TL-0042-8817", "Existing term loan", C(ws, 22000, None, INR), 7, C(ws, date(2027, 6, 7), None, DT)])
ws = sheet("Lenders", ["lender", "type", "annual_rate", "tenure_months", "min_amount_inr", "max_amount_inr", "min_gst_months",
                       "offer_valid", "processing_fee_pct", "requested_inr", "gst_months_available", "match", "monthly_instalment_inr",
                       "total_interest_inr", "why", "emis_after_inr", "credit_cap_inr", "within_cap"],
           [9, 6, 10, 12, 13, 13, 13, 10, 15, 12, 18, 7, 20, 16, 26, 14, 14, 10], formula_cols=tuple(range(9, 18)))
for i, l in enumerate(D["LENDERS"], start=2):
    ws.append([l[0], l[1], C(ws, l[2], None, PCT), l[3], C(ws, l[4], None, INR), C(ws, l[5], None, INR), l[6], l[7], C(ws, l[8], None, PCT),
               C(ws, f"={a('Drawdown amount (option 2)')}", None, INR), f"=COUNTA(GST_Returns!$A$2:$A${NG})",
               f"=AND(H{i},J{i}>=E{i},J{i}<=F{i},K{i}>=G{i})",
               C(ws, f'=IF(L{i},ROUND(PMT(C{i}/12,D{i},-J{i}),-1),"")', None, INR),
               C(ws, f'=IF(L{i},M{i}*D{i}-J{i},"")', None, INR),
               f'=IF(L{i},"Match",IF(J{i}<E{i},"Minimum ₹"&TEXT(E{i},"#,##0"),IF(K{i}<G{i},"Needs "&G{i}&" months of GST",IF(J{i}>F{i},"Above maximum","Offer not valid"))))',
               C(ws, f'=IF(L{i},M{i}+SUM(Existing_EMIs!$C$2:$C$2),"")', None, INR),
               C(ws, "=Calibration!$C$12", None, INR), f'=IF(L{i},P{i}<=Q{i},"")'])
ws = sheet("Owner_Mandate", ["rule_key", "value", "unit", "guard_rule"], [38, 70, 8, 10])
for m in D["MANDATE"]: ws.append([m[0], C(ws, m[1], F_IN), m[2], m[3]])

# ------------------------------------------------------------------ Monthly_Summary
ws = sheet("Monthly_Summary", ["month_start", "month_end", "outlet_id", "paytm_net_inr", "all_channel_collections_inr", "paytm_share",
                               "other_qr_inr", "cash_incl_till_expenses_inr"], [12, 12, 9, 15, 24, 11, 14, 22], formula_cols=(1, 3, 4, 5, 6, 7))
i = 2
y, mth = 2025, 5
while (y, mth) <= (2026, 10):
    for o in ("DAD", "MAT", "SIO"):
        ws.append([C(ws, date(y, mth, 1), None, DT), C(ws, f"=EOMONTH(A{i},0)", None, DT), o,
                   C(ws, f'=SUMIFS({PT("J")},{PT("G")},C{i},{PT("C")},">="&A{i},{PT("C")},"<="&B{i})', None, INR),
                   C(ws, f'=SUMIFS({SH("M")},{SH("C")},C{i},{SH("B")},">="&A{i},{SH("B")},"<="&B{i})', None, INR),
                   C(ws, f'=IF(E{i}=0,0,SUMIFS({SH("G")},{SH("C")},C{i},{SH("B")},">="&A{i},{SH("B")},"<="&B{i})/E{i})', None, PCT),
                   C(ws, f'=SUMIFS({SH("H")},{SH("C")},C{i},{SH("B")},">="&A{i},{SH("B")},"<="&B{i})+SUMIFS({SH("I")},{SH("C")},C{i},{SH("B")},">="&A{i},{SH("B")},"<="&B{i})', None, INR),
                   C(ws, f'=SUMIFS({SH("J")},{SH("C")},C{i},{SH("B")},">="&A{i},{SH("B")},"<="&B{i})+SUMIFS({SH("K")},{SH("C")},C{i},{SH("B")},">="&A{i},{SH("B")},"<="&B{i})', None, INR)])
        i += 1
    mth += 1
    if mth == 13: y, mth = y + 1, 1
ws.append([C(ws, "Note: October 2026 runs to 30 Oct; Paytm includes today's partial morning, handovers end 29 Oct.", F_NOTE)])

# ------------------------------------------------------------------ Recurring_Bills (derived from statement narrations, never from label_*)
ws = sheet("Recurring_Bills", ["stream", "narration_key", "already_in_upcoming_outflows", "occurrences_since_1_may_2026",
                               "amount_per_month_inr", "typical_day", "next_due_date", "include_in_cash_plan"],
           [26, 22, 26, 26, 20, 11, 14, 18], formula_cols=(3, 4, 5, 6, 7))
REC = [("Existing term loan EMI", "TL EMI", False), ("Electricity Dadar", "ELECTRICITY/CA1000", False),
       ("Electricity Matunga", "ELECTRICITY/CA1001", False), ("Electricity Sion", "ELECTRICITY/CA1002", False),
       ("Broadband, 3 outlets", "BROADBAND", False), ("Rent, 3 outlets", "/RENT ", True), ("Salaries", "NEFT/SAL ", True),
       ("GST payment", "GST CHALLAN", True)]
for i, (nm, key, known) in enumerate(REC, start=2):
    hit = f'ISNUMBER(SEARCH(B{i},{BS("H")}))*({BS("G")}>=DATE(2026,5,1))'
    ws.append([nm, C(ws, key, F_IN), C(ws, known, F_IN),
               f'=COUNTIFS({BS("H")},"*"&B{i}&"*",{BS("G")},">="&DATE(2026,5,1))',
               C(ws, f'=ROUND(SUMIFS({BS("D")},{BS("H")},"*"&B{i}&"*",{BS("G")},">="&DATE(2026,8,1),{BS("G")},"<="&DATE(2026,10,31))/3,0)', None, INR),
               f"=IF(D{i}=0,\"\",ROUND(SUMPRODUCT({hit}*DAY({BS('G')}))/D{i},0))",
               C(ws, f'=IF(D{i}=0,"",DATE(2026,11,F{i}))', None, DT),
               f"=AND(D{i}>=5,NOT(C{i}))"])
NREC = len(REC) + 1
ws.append([])
ws.append([C(ws, "Recurring streams are found by matching narration text in Bank_Statement (case-insensitive). "
               "Amount = average of Aug-Oct 2026. Streams already in Upcoming_Outflows are excluded to avoid double counting. "
               "Production: confirm against BBPS / UPI AutoPay mandates where available.", F_NOTE)])

# ------------------------------------------------------------------ Cash_Plan
ws = wb.create_sheet("Cash_Plan")
for col, w in zip("ABCDEFGH", (14, 18, 16, 18, 16, 18, 22, 40)): ws.column_dimensions[col].width = w
ws.append([C(ws, "Cash plan · 30 Oct to 20 Nov 2026 (cash + bank combined)", F_T)])
ws.append(["Opening, start of 30 Oct", C(ws, f"=ROUND({LAST_BAL}+{a('Cash in hand across tills, close of 29 Oct 2026')},0)", F_B, INR),
           C(ws, "Bank close 29 Oct + cash in tills", F_NOTE)])
ws.append([])
ws.append([C(ws, h, F_HDR, fill=FILL_F) for h in ("date", "inflow_forecast_inr", "known_outflows_inr", "recurring_bills_inr",
                                                   "closing_inr", "drawdown_option2_inr", "closing_with_option2_inr")])
r0 = 5
for k in range(22):
    i = r0 + k; d = date(2026, 10, 30) + timedelta(k)
    prevE = "$B$2" if k == 0 else f"E{i - 1}"
    prevG = "$B$2" if k == 0 else f"G{i - 1}"
    ws.append([C(ws, d, None, DT),
               C(ws, f"=IF(A{i}<{a('Forecast step-up date')},{a('Forecast combined inflow a day, 30 Oct to 3 Nov')},{a('Forecast combined inflow a day, from 4 Nov')})", None, INR),
               C(ws, f"=SUMIFS(Upcoming_Outflows!$D$2:$D$8,Upcoming_Outflows!$B$2:$B$8,A{i})", None, INR),
               C(ws, f"=SUMIFS(Recurring_Bills!$E$2:$E${NREC},Recurring_Bills!$G$2:$G${NREC},A{i},Recurring_Bills!$H$2:$H${NREC},TRUE)", None, INR),
               C(ws, f"={prevE}+B{i}-C{i}-D{i}", None, INR),
               C(ws, f"=IF(A{i}={a('Drawdown date')},{a('Drawdown amount (option 2)')},0)", None, INR),
               C(ws, f"={prevG}+B{i}-C{i}-D{i}+F{i}", None, INR)])
rL = r0 + 21
ws.append([])
summ = [
    ("Lowest closing", f"=MIN(E{r0}:E{rL})", INR),
    ("Date of lowest", f"=INDEX(A{r0}:A{rL},MATCH(MIN(E{r0}:E{rL}),E{r0}:E{rL},0))", DT),
    ("First day back above zero", f"=INDEX(A{r0}:A{rL},MATCH(1,INDEX((E{r0}:E{rL}>0)*(A{r0}:A{rL}>B{rL + 3}),0),0))", DT),
    ("Lowest closing with option 2", f"=MIN(G{r0}:G{rL})", INR),
    ("Closing on 5 Nov with option 2", f"=INDEX(G{r0}:G{rL},MATCH(DATE(2026,11,5),A{r0}:A{rL},0))", INR),
    ("Option 2 interest if repaid by 30 Nov (approx.)",
     f"=ROUND({a('Drawdown amount (option 2)')}*{a('Drawdown interest rate (Lender A)')}*({a('Early repayment date')}-{a('Drawdown date')})/365,-2)", INR),
    ("Drawdown counted once as inflow; GD-1182 outflow kept at full ₹5,20,000", "", None),
]
for lab, f, fmt in summ:
    ws.append([lab, C(ws, f, F_B, fmt) if f else None])

# ------------------------------------------------------------------ Calibration
ws = wb.create_sheet("Calibration")
for col, w in zip("ABCDEF", (6, 52, 16, 14, 12, 10)): ws.column_dimensions[col].width = w
ws.append([C(ws, "Calibration · PRD 01 targets, computed from the data", F_T)])
ws.append([C(ws, h, F_HDR, fill=FILL_F) for h in ("#", "Check", "Measured", "Target", "Tolerance", "Result")])
W = f'{SH("B")},">="&{a("Calibration window from")},{SH("B")},"<="&{a("Calibration window to")}'
def coll(o, frm, to):
    return f'SUMIFS({SH("M")},{SH("C")},"{o}",{SH("B")},">="&{frm},{SH("B")},"<="&{to})'
rows = []
for o, t in (("DAD", 900000), ("MAT", 600000), ("SIO", 400000)):
    rows.append((f"Average monthly collections {o}, 12 months", f'=SUMIFS({SH("M")},{SH("C")},"{o}",{W})/12', t, f"={a('Monthly sales tolerance')}", "rel", INR))
rows.append(("Average monthly collections, all outlets", f'=SUMIFS({SH("M")},{W})/12', 1900000, f"={a('Monthly sales tolerance')}", "rel", INR))
rows.append(("Paytm share of collections", f'=SUMIFS({SH("G")},{W})/SUMIFS({SH("M")},{W})', f"={a('Paytm share target')}", f"={a('Paytm share tolerance (points)')}", "abs", PCT))
for o, t in (("DAD", 1.8), ("MAT", 1.5), ("SIO", 1.5)):
    rows.append((f"Diwali week 2025 vs normal week, {o}",
                 f"={coll(o, a('Diwali week 2025 from'), a('Diwali week 2025 to'))}/({coll(o, a('Reference weeks from (4 normal weeks)'), a('Reference weeks to'))}/4)",
                 t, f"={a('Diwali uplift tolerance (×)')}", "abs", "0.00\"×\""))
rows.append(("Opening balance, start of 30 Oct (bank + tills)", "=Cash_Plan!$B$2", 450000, 0, "abs", INR))
rows.append(("Credit cap (8% of average monthly collections)", f"=ROUND({a('Credit cap, % of average monthly collections')}*C6,0)", 152000, 0, "abs", INR))
for o, t in (("DAD", 4980), ("MAT", 3120), ("SIO", 2040)):
    rows.append((f"Paytm today (30 Oct) {o}", f'=SUMIFS({PT("J")},{PT("G")},"{o}",{PT("C")},{a("Today")})', t, 0, "abs", INR))
thu = lambda frm, to: f'(({SH("C")}="MAT")*({SH("D")}="Evening")*(WEEKDAY({SH("B")})=5)*({SH("B")}>={frm})*({SH("B")}<={to}))'
rows.append(("Matunga Thursday evening expected (avg 6 Aug to 1 Oct)",
             f'=SUMPRODUCT({thu("DATE(2026,8,6)", "DATE(2026,10,1)")}*{SH("M")})/SUMPRODUCT({thu("DATE(2026,8,6)", "DATE(2026,10,1)")}*1)',
             17600, 0.02, "rel", INR))
rows.append(("Matunga Thursday evenings below range, last 4 (8-29 Oct)",
             f'=SUMPRODUCT({thu("DATE(2026,10,8)", "DATE(2026,10,29)")}*({SH("M")}<C16*(1-{a("Shift range band (±)")})))', 3, 0, "abs", "0"))
rows.append(("Known outflows 31 Oct to 20 Nov", "=Upcoming_Outflows!$D$10", 1266500, 0, "abs", INR))
rows.append(("Cash plan lowest point", "=Cash_Plan!$B$28", -184500, 0, "abs", INR))
rows.append(("Closing on 5 Nov with ₹2 lakh drawdown", "=Cash_Plan!$B$32", 15500, 0, "abs", INR))
rows.append(("Lender A monthly instalment", "=Lenders!$M$2", 68730, 0, "abs", INR))
rows.append(("EMIs after drawdown (Lender A)", "=Lenders!$P$2", 90730, 0, "abs", INR))
rows.append(("Paytm settlement days not matching the bank", f'=COUNTIF(Paytm_Settlements!$I$2:$I${NSET},"CHECK")', 0, 0, "abs", "0"))
for n, (lab, f, t, tol, kind, fmt) in enumerate(rows, start=1):
    i = n + 2
    res = (f'=IF(ABS(C{i}/D{i}-1)<=E{i},"PASS","FAIL")' if kind == "rel" else f'=IF(ABS(C{i}-D{i})<=E{i},"PASS","FAIL")')
    ws.append([n, lab, C(ws, f, None, fmt), C(ws, t, F_IN, fmt), C(ws, tol, None, PCT if kind == "rel" or fmt == PCT else "0.00"), res])
ws.append([])
ws.append([None, C(ws, "Checks passed", F_B), C(ws, f'=COUNTIF(F3:F{len(rows) + 2},"PASS")&" of {len(rows)}"', F_B)])

# ------------------------------------------------------------------ Invoice bodies (invoices.py)
if "INV_LINES" in D:
    L = D["INV_LINES"]; NL = len(L) + 1
    ws = sheet("Invoice_Lines", ["invoice_row", "invoice_no", "line_no", "description", "hsn", "unit", "qty", "rate_inr", "discount_pct",
                                 "taxable_value_inr", "gst_rate", "cgst_inr", "sgst_inr", "line_total_inr"],
               [10, 10, 7, 58, 7, 6, 6, 10, 11, 15, 9, 11, 11, 13], formula_cols=(9, 11, 12, 13))
    for i, l in enumerate(L, start=2):
        ws.append([l["invoice_row"], l["invoice_no"], l["line_no"], l["description"], l["hsn"], l["unit"], l["qty"],
                   C(ws, l["rate_inr"], None, INR2), C(ws, l["discount_pct"], None, PCT),
                   C(ws, f"=ROUND(G{i}*H{i}*(1-I{i}),2)", None, INR2), C(ws, l["gst_rate"], None, PCT),
                   C(ws, f"=ROUND(J{i}*K{i}/2,2)", None, INR2), C(ws, f"=ROUND(J{i}*K{i},2)-L{i}", None, INR2),
                   C(ws, f"=J{i}+L{i}+M{i}", None, INR2)])
    IL = lambda col: f"Invoice_Lines!${col}$2:${col}${NL}"
    Hd = D["INV_HEADERS"]
    cols = ["invoice_row", "invoice_no", "supplier_name", "supplier_gstin", "issue_date", "due_date", "bank_account_masked",
            "bank_ifsc", "received_via", "scenario", "notes", "round_off", "invoice_total", "lines", "taxable_value_inr", "gst_inr",
            "computed_total_inr", "ties_to_header", "earlier_invoice_with_same_supplier_total_and_lines", "body_status",
            "evidence_check"]
    ws = sheet("Invoice_Headers", cols, [10, 10, 24, 17, 11, 11, 16, 12, 22, 15, 60, 9, 13, 7, 16, 12, 17, 13, 30, 13, 40],
               formula_cols=(13, 14, 15, 16, 17, 18, 20))
    for i, h in enumerate(Hd, start=2):
        ws.append([h["invoice_row"], h["invoice_no"], h["supplier_name"], h["supplier_gstin"],
                   C(ws, date.fromisoformat(h["issue_date"]), None, DT), C(ws, date.fromisoformat(h["due_date"]), None, DT),
                   h["bank_account_masked"], h["bank_ifsc"], h["received_via"], h["scenario"], h["notes"],
                   C(ws, h["round_off"], None, INR2), C(ws, h["invoice_total"], None, INR),
                   f"=COUNTIF({IL('A')},A{i})", C(ws, f"=SUMIFS({IL('J')},{IL('A')},A{i})", None, INR2),
                   C(ws, f"=SUMIFS({IL('L')},{IL('A')},A{i})+SUMIFS({IL('M')},{IL('A')},A{i})", None, INR2),
                   C(ws, f"=ROUND(O{i}+P{i}+L{i},2)", None, INR2), f'=IF(ABS(Q{i}-M{i})<0.01,"YES","NO")',
                   # R12 duplicate signal: another invoice (different row) with the same GSTIN, same total and same line count
                   f'=IF(COUNTIFS($D$2:$D${len(Hd) + 1},D{i},$M$2:$M${len(Hd) + 1},M{i},$N$2:$N${len(Hd) + 1},N{i},$E$2:$E${len(Hd) + 1},"<"&E{i},$B$2:$B${len(Hd) + 1},"<>"&B{i},$N$2:$N${len(Hd) + 1},">0")>0,"YES","NO")',
                   h["body_status"],
                   f'=IF(N{i}=0,"NEEDS_VERIFICATION (body missing)",IF(R{i}="NO","NEEDS_VERIFICATION (lines do not tie)",'
                   f'IFERROR(IF(G{i}<>"XXXXXXXX"&INDEX(Suppliers!$D$2:$D$8,MATCH(C{i},Suppliers!$B$2:$B$8,0)),"BLOCK (account not verified)",'
                   f'IF(S{i}="YES","CHECK DUPLICATE (R12)","EVIDENCE OK (PIN still needed)")),"BLOCK (payee not on approved list)")))'])

if "UPLOADS" in D:
    U = D["UPLOADS"]
    ws = sheet("Bill_Uploads", list(U[0]), [9, 24, 28, 10, 10, 15, 12, 12, 52, 44, 70, 20])
    for u in U:
        ws.append(list(u.values()))

wb.save(f"{OUT}/PayTMunim_Synthetic_Data.xlsx")
print("saved", NP, NB, NH)
